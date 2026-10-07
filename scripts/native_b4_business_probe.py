"""Exercise local review/export on the three real native B3 analyses; no model calls."""
import argparse
import hashlib
import json
from io import BytesIO
from pathlib import Path
import sqlite3
import urllib.request
from openpyxl import load_workbook
from backend.invoice_xlsx import COLUMNS
from decimal import Decimal

ROOT = Path(__file__).resolve().parents[1]
REPORT = ROOT / 'logs&architechture/AP_native_agent_B4_verification.json'
BASE = 'http://127.0.0.1:8000'


def api(path, body=None, method=None):
    request = urllib.request.Request(BASE+path, data=None if body is None else json.dumps(body).encode(),
        method=method, headers={'Content-Type':'application/json','X-AP-Review':'1'})
    with urllib.request.urlopen(request, timeout=30) as response:
        return json.load(response)


def main():
    parser=argparse.ArgumentParser()
    parser.add_argument('action',choices=['pending','decide','export','snapshot','verify'])
    action=parser.parse_args().action
    samples=json.loads((ROOT/'logs&architechture/AP_native_agent_B3_verification.json').read_text(encoding='utf-8'))['samples']
    report=json.loads(REPORT.read_text(encoding='utf-8')) if REPORT.exists() else {}
    states={s['case']:api('/api/ap/invoices/'+s['invoice_id']+'/review') for s in samples}
    register=api('/api/ap/export-register?only_unexported=false')['rows']
    if action=='pending':
        assert states['scan']['status']==states['missing']['status']=='pending'
        assert states['text']['status']=='approved' and states['text']['approval_method']=='automatic'
        ids={r['invoice_id'] for r in register}
        assert all(s['invoice_id'] not in ids for s in samples if s['case']!='text')
        report['pending_excluded']=True
    elif action=='decide':
        for s in samples:
            if s['case']=='text':continue
            state=states[s['case']];path='/api/ap/invoices/'+s['invoice_id']+'/review'
            if state['status']=='pending':
                if s['case']=='scan':
                    state=api(path,{'expected_revision':state['current_revision'],
                        'processing_attempt_id':state['processing_attempt_id'],
                        'changes':{'supplier_address':'12 Sample Street\nBerlin 10115'},'page':1,
                        'reason':'B4 synthetic demo: preserve the address and format its city on a separate line.'},'PATCH')
                elif state['current_revision']==0:
                    state=api(path,{'expected_revision':state['current_revision'],
                        'processing_attempt_id':state['processing_attempt_id'],'changes':{},'page':1,
                        'reason':'B4 synthetic demo: inspected page 1; due date is absent.'},'PATCH')
                state=api(path+'/decision',{'expected_revision':state['current_revision'],
                    'processing_attempt_id':state['processing_attempt_id'],
                    'decision':'approve' if s['case']=='scan' else 'reject',
                    'request_id':'native-b4-20261007-'+s['case'],
                    'reason':'Synthetic scan checked against page 1.' if s['case']=='scan' else 'Synthetic invoice has no due date; request corrected document.'},'POST')
            assert state['status']==('approved' if s['case']=='scan' else 'rejected')
            report.setdefault('decisions',{})[s['case']]={k:state[k] for k in ('status','current_revision','processing_attempt_id','approval_method','event')}
    elif action=='export':
        events=[states[k]['event']['event_id'] for k in ('text','scan')]
        report['batch']=api('/api/ap/export-batches',{'event_ids':events,'only_unexported':False},'POST')
    else:
        assert [states[k]['status'] for k in ('text','scan','missing')]==['approved','approved','rejected']
        assert states['scan']['current_fields']['supplier_address']=='12 Sample Street\nBerlin 10115'
        assert states['scan']['raw_extraction']['fields']['supplier_address']=='12 Sample Street, Berlin 10115'
        assert states['scan']['approval_method']=='manual'
        missing=next(s['invoice_id'] for s in samples if s['case']=='missing')
        assert missing not in {r['invoice_id'] for r in register}
        batch=api('/api/ap/export-batches/'+report['batch']['export_id'])
        notices=[api('/api/ap/result-notifications/review/'+s['event']['event_id']) for s in states.values()]
        notices.append(api('/api/ap/result-notifications/batch/'+batch['export_id']))
        assert all(n['status']=='sent' for n in notices)
        assert all(s['event']['status']=='processed' for s in states.values())
        data=urllib.request.urlopen(BASE+'/api/ap/export-batches/'+batch['export_id']+'/files/xlsx').read()
        assert hashlib.sha256(data).hexdigest()==batch['file_sha256']
        book=load_workbook(BytesIO(data),data_only=False);sheet=book.active
        headers=[c.value for c in sheet[1]]
        rows=[dict(zip(headers,values)) for values in sheet.iter_rows(min_row=2,max_row=3,values_only=True)]
        assert {r['Invoice number'] for r in rows}=={'NATIVE-B3-20261007-TEXT','NATIVE-B3-20261007-SCAN'}
        assert sum(r['Total amount'] for r in rows)==12019
        assert {r['Approval method'] for r in rows}=={'automatic','manual'}
        assert all(r['Invoice ID']!=missing for r in rows)
        for row in rows:
            state=states['text' if row['Invoice number'].endswith('TEXT') else 'scan']
            assert row['Review revision']==state['current_revision']
            for key,label,_ in COLUMNS:
                if key not in state['current_fields']:continue
                expected=state['current_fields'][key];actual=row[label]
                if expected is not None and key.endswith('_amount'):
                    assert Decimal(str(actual))==Decimal(expected)
                elif expected is not None and key in ('invoice_date','due_date'):
                    assert actual.date().isoformat()==expected
                else:assert actual==expected,(key,actual,expected)
        book.close()
        db=sqlite3.connect((ROOT/'data/ap/invoices.sqlite3').as_uri()+'?mode=ro',uri=True)
        assert db.execute('PRAGMA integrity_check').fetchone()[0]=='ok'
        assert not db.execute('PRAGMA foreign_key_check').fetchall()
        assert db.execute('SELECT COUNT(*) FROM processing_attempts').fetchone()[0]==103
        snapshot={'events':[s['event'] for s in states.values()],'notices':notices,'batch':batch,
            'counts':{t[0]:db.execute('SELECT COUNT(*) FROM "'+t[0]+'"').fetchone()[0] for t in db.execute("SELECT name FROM sqlite_master WHERE type='table' AND name NOT LIKE 'sqlite_%'")}}
        if action=='verify':
            assert report['snapshot']==snapshot,'Persistent business state changed after restart/replay'
            report['restart_replay_unchanged']=True
        else:report['snapshot']=snapshot
        report['excel']={'sha256':batch['file_sha256'],'row_count':2,'total_eur':12019,'rejected_excluded':True,'approved_versions_preserved':True}
    REPORT.write_text(json.dumps(report,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'action':action,'passed':True}))


if __name__=='__main__':main()
