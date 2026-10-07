"""Register filtering, immutable batch membership, concurrency and recovery without model calls."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import date
from io import BytesIO
import json
import sqlite3
import subprocess
import sys
from uuid import uuid4

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from backend.app import create_app
from backend.exports import Exports
from backend.invoice_xlsx import NOTICE
from backend.migrations import migrate
from backend.test_exports import approved, export
from backend.test_review_decisions import env, seed, save, decide, HEADERS


def batch(env, events, **extra):
    return env[0].post('/api/ap/export-batches', json={'event_ids': events, **extra}, headers=HEADERS)


def preview(env, **params):
    response = env[0].get('/api/ap/export-register', params=params)
    assert response.status_code == 200, response.text
    return response.json()


def test_register_appears_at_approval_without_file_or_event_delivery(env):
    invoice, _ = seed(env)
    assert preview(env)['count'] == 0
    save(env, invoice)
    decide(env, invoice)
    decide(env, invoice)
    automatic, _, _ = approved(env, automatic=True)
    rejected, _ = seed(env)
    save(env, rejected)
    decide(env, rejected, 'reject')
    assert {r['invoice_id'] for r in preview(env)['rows']} == {invoice, automatic}
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*) FROM export_batches').fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM exports').fetchone()[0] == 0
        assert db.execute('SELECT COUNT(*) FROM review_events WHERE attempts!=0').fetchone()[0] == 0
    assert not (env[2] / 'exports').exists()


def test_filters_are_inclusive_literal_and_currency_totals_are_separate(env):
    approved(env, total='0.10', changes={'supplier_name': 'ACME % Ltd'})
    approved(env, total='0.20', changes={'supplier_name': 'ACME % Ltd', 'invoice_date': '2026-10-02'})
    approved(env, total='200.00', changes={'supplier_name': 'Other', 'currency': 'USD'})
    all_rows = preview(env)
    assert all_rows['totals_by_currency'] == {'EUR': '0.30', 'USD': '200.00'}
    result = preview(env, supplier='acme %', date_from='2026-10-02', date_to='2026-10-02', currency='eur')
    assert result['count'] == 1 and result['totals_by_currency'] == {'EUR': '0.20'}
    assert preview(env, supplier='<script>')['count'] == 0
    assert env[0].get('/api/ap/export-register?date_from=2026-10-03&date_to=2026-10-01').status_code == 422
    assert env[0].get('/api/ap/export-register?date_to=bad-date').status_code == 422


def test_combined_rows_are_exact_snapshots_and_download_is_read_only(env):
    invoice, attempt, first = approved(env, automatic=True, number='0000123', changes={'supplier_name': '=1+1', 'supplier_tax_id': '00123'})
    _, _, second = approved(env, changes={'currency': 'KWD'}, total='1.25')
    # A later machine result cannot replace the immutable approval.
    with closing(env[1]()) as db:
        db.execute('UPDATE processing_attempts SET result_json=? WHERE attempt_id=?', (json.dumps({'fields': {'total_amount': '999'}}), attempt))
        db.commit()
    result = batch(env, [first, second])
    assert result.status_code == 200, result.text
    record = result.json()
    assert record['row_count'] == 2 and record['attempts'] == 1
    workbook = load_workbook(BytesIO(env[0].get(record['download_url']).content))
    sheet = workbook.active
    assert sheet.auto_filter.ref == 'A1:S3' and sheet['A5'].value == NOTICE
    values = {sheet.cell(row, 1).value: row for row in (2, 3)}
    row = values[invoice]
    assert sheet.cell(row, 3).value == '=1+1' and sheet.cell(row, 3).data_type == 's'
    assert sheet.cell(row, 5).value == '00123' and sheet.cell(row, 6).value == '0000123'
    assert sheet.cell(row, 12).value == 119 and sheet.cell(row, 13).value is None
    assert sheet.cell(row, 7).value.date() == date(2026, 10, 1)
    assert sheet.cell(row, 18).value == 'automatic'
    assert sheet.cell(5 - row, 12).number_format == '#,##0.000 "KWD"'
    assert all(sheet.cell(r, 16).value == record['export_id'] for r in (2, 3))
    assert not any(c.data_type == 'f' or c.hyperlink for cells in sheet for c in cells)
    workbook.close()
    assert batch(env, [second, first], only_unexported=False).json() == record
    assert preview(env)['count'] == 0
    assert preview(env, only_unexported=False)['count'] == 2
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*),MAX(attempts) FROM export_batches').fetchone()[:] == (1, 1)
        assert db.execute('SELECT COUNT(*) FROM exports').fetchone()[0] == 0
        assert db.execute('PRAGMA foreign_key_check').fetchall() == []


def test_incremental_selection_is_frozen_and_legacy_exports_count(env):
    _, _, legacy = approved(env)
    export(env, legacy)
    _, _, first = approved(env)
    shown = [r['event_id'] for r in preview(env)['rows']]
    assert shown == [first]
    _, _, second = approved(env)
    record = batch(env, shown).json()
    assert record['row_count'] == 1
    assert [r['event_id'] for r in preview(env)['rows']] == [second]
    assert batch(env, [legacy, second]).status_code == 409
    result = batch(env, [legacy, first, second], only_unexported=False)
    assert result.status_code == 200 and result.json()['row_count'] == 3


def test_invalid_selection_is_atomic_and_local_write_protection_applies(env):
    _, _, valid = approved(env)
    invoice, _ = seed(env)
    save(env, invoice)
    rejected = decide(env, invoice, 'reject').json()['event']['event_id']
    assert batch(env, [valid, rejected]).status_code == 409
    assert batch(env, [valid, str(uuid4())]).status_code == 404
    assert batch(env, [valid, valid]).status_code == 422
    assert batch(env, []).status_code == 422
    assert batch(env, [valid] * 501).status_code == 422
    assert batch(env, [valid], file_path='../bad.xlsx').status_code == 422
    assert env[0].post('/api/ap/export-batches', json={'event_ids': [valid]}).status_code == 403
    assert env[0].post('/api/ap/export-batches', json={'event_ids': [valid]}, headers={**HEADERS, 'Origin': 'https://evil.invalid'}).status_code == 403
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*) FROM export_batches').fetchone()[0] == 0


def test_same_selection_parallel_and_mapping_change_reuses_identity(env, monkeypatch):
    events = [approved(env)[2], approved(env)[2]]
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda _: batch(env, events), range(4)))
    assert all(r.status_code in (200, 409) for r in responses)
    first = batch(env, events).json()
    monkeypatch.setattr('backend.export_register.MAPPING_VERSION', 'future-mapping')
    assert batch(env, events[::-1]).json() == first
    assert len(list((env[2] / 'exports/batches').glob('*.xlsx'))) == 1


def test_overlapping_unexported_batches_cannot_both_succeed(env):
    events = [approved(env)[2] for _ in range(3)]
    with ThreadPoolExecutor(max_workers=2) as pool:
        responses = list(pool.map(lambda items: batch(env, items), [events[:2], events[1:]]))
    assert sorted(r.status_code for r in responses) == [200, 409]
    assert preview(env)['count'] == 1


def test_batch_membership_and_identity_are_immutable(env):
    record = batch(env, [approved(env)[2]]).json()
    with closing(env[1]()) as db:
        for sql in ['DELETE FROM export_batch_items', 'UPDATE export_batch_items SET ordinal=42',
                    "UPDATE export_batches SET selection_hash='changed'", "UPDATE export_batches SET mapping_version='changed'"]:
            with pytest.raises(sqlite3.IntegrityError):
                db.execute(sql)
            db.rollback()
    assert env[0].get(record['download_url']).status_code == 200


@pytest.mark.parametrize('damage', ['missing', 'corrupt'])
def test_batch_recovery_after_restart_is_explicit_and_preserves_hash(env, damage):
    record = batch(env, [approved(env)[2], approved(env)[2]]).json()
    path = env[2] / record['file_path']
    if damage == 'missing':
        path.unlink()
    else:
        path.write_bytes(b'not a workbook')
    with TestClient(create_app(env[2]), base_url='http://localhost') as restarted:
        assert restarted.get(record['download_url']).status_code == 409
        assert restarted.get('/api/ap/export-batches/' + record['export_id']).json()['file_health'] == 'missing_or_invalid'
        assert restarted.post('/api/ap/export-batches/' + record['export_id'] + '/retry').status_code == 403
        recovered = restarted.post('/api/ap/export-batches/' + record['export_id'] + '/retry', headers=HEADERS).json()
    assert recovered['export_id'] == record['export_id'] and recovered['file_sha256'] == record['file_sha256']
    assert recovered['members'] == record['members'] and recovered['attempts'] == 2


def test_batch_commit_failure_reuses_file_and_visible_history(env, monkeypatch):
    events = [approved(env)[2], approved(env)[2]]
    original = Exports.complete
    def fail(*args):
        raise sqlite3.OperationalError('injected')
    monkeypatch.setattr(Exports, 'complete', fail)
    assert batch(env, events).status_code == 503
    saved = env[0].get('/api/ap/export-batches').json()[0]
    assert saved['status'] == 'failed' and saved['row_count'] == 2
    assert preview(env)['count'] == 2
    monkeypatch.setattr(Exports, 'complete', original)
    monkeypatch.setattr(Exports, 'write_file', lambda *args: pytest.fail('Reuse verified orphan'))
    recovered = batch(env, events).json()
    assert recovered['export_id'] == saved['export_id'] and recovered['attempts'] == 2


def test_batch_write_attempts_are_bounded_and_keep_approval(env, monkeypatch):
    events = [approved(env)[2]]
    def fail(*args):
        raise PermissionError('private-path')
    monkeypatch.setattr(Exports, 'write_file', fail)
    for _ in range(5):
        response = batch(env, events)
        assert response.status_code == 503 and 'private-path' not in response.text
    assert batch(env, events).status_code == 409
    assert preview(env)['count'] == 1


def test_failed_repair_does_not_make_previously_exported_invoice_new(env, monkeypatch):
    events = [approved(env)[2]]
    saved = batch(env, events).json()
    (env[2] / saved['file_path']).unlink()
    def fail(*args):
        raise PermissionError('injected')
    monkeypatch.setattr(Exports, 'write_file', fail)
    assert batch(env, events).status_code == 503
    assert preview(env)['count'] == 0
    assert preview(env, only_unexported=False)['rows'][0]['exported'] is True
    with closing(env[1]()) as db:
        assert db.execute('SELECT status,file_sha256 FROM export_batches').fetchone()[:] == ('failed', saved['file_sha256'])


def test_stale_approval_or_changed_manifest_blocks_download(env):
    invoice, _, event = approved(env)
    record = batch(env, [event]).json()
    with closing(env[1]()) as db:
        db.execute('UPDATE review_tasks SET current_revision=current_revision+1 WHERE invoice_id=?', (invoice,))
        db.commit()
    assert env[0].get(record['download_url']).status_code == 409
    assert batch(env, [event]).status_code == 409


def test_migration_three_to_four_keeps_single_exports_and_repeats(env):
    legacy = export(env, approved(env)[2]).json()
    with closing(env[1]()) as db:
        # Recreate the actual version-three boundary without altering business data.
        db.execute('DROP TABLE outcome_notification_attempts')
        db.execute('DROP TABLE outcome_notifications')
        db.execute('DROP TABLE export_batch_items')
        db.execute('DROP TABLE export_batches')
        db.execute('PRAGMA user_version=3')
        db.commit()
        before = [tuple(r) for r in db.execute('SELECT * FROM exports')]
        assert migrate(db) == 5 and migrate(db) == 5
        assert [tuple(r) for r in db.execute('SELECT * FROM exports')] == before
        assert db.execute('PRAGMA foreign_key_check').fetchall() == []
    assert env[0].get(legacy['download_url']).status_code == 200
    assert preview(env)['count'] == 0 and preview(env, only_unexported=False)['count'] == 1


@pytest.mark.parametrize('phase', ['partial', 'renamed'])
def test_batch_process_exit_keeps_identity_and_membership(env, phase):
    events = [approved(env)[2], approved(env)[2]]
    script = '''
import os,sqlite3,sys,json
from pathlib import Path
from backend.export_register import ExportBatches,BatchRequest
root=Path(sys.argv[1])
def connect():
 db=sqlite3.connect(root/'invoices.sqlite3',timeout=30);db.row_factory=sqlite3.Row;db.execute('PRAGMA foreign_keys=ON');return db
class CrashBatch(ExportBatches):
 def write_file(self,path,data):
  if sys.argv[3]=='partial':
   path.parent.mkdir(parents=True,exist_ok=True);path.with_suffix('.tmp').write_bytes(data[:50])
  else:
   super().write_file(path,data)
  os._exit(77)
CrashBatch(connect,root).create_batch(BatchRequest(event_ids=json.loads(sys.argv[2])))
'''
    result = subprocess.run([sys.executable, '-c', script, str(env[2]), json.dumps(events), phase], capture_output=True, timeout=30)
    assert result.returncode == 77, result.stderr
    assert batch(env, events).status_code == 409
    with closing(env[1]()) as db:
        saved = dict(db.execute('SELECT * FROM export_batches').fetchone())
        assert saved['status'] == 'writing' and saved['attempts'] == 1
        db.execute("UPDATE export_batches SET updated_at='2000-01-01T00:00:00+00:00'")
        db.commit()
    recovered = batch(env, events).json()
    assert recovered['export_id'] == saved['export_id'] and recovered['row_count'] == 2 and recovered['attempts'] == 2
    assert not list((env[2] / 'exports/batches').glob('*.tmp'))
