"""Versioned invoice workbook mapping; all business values come from a saved snapshot."""
from datetime import date, datetime
from decimal import Decimal
import hashlib
from io import BytesIO
from zipfile import ZipFile, ZipInfo, ZIP_DEFLATED

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from backend.validation import CURRENCY_PRECISION, MONEY_FIELDS

MAPPING_VERSION = 'invoice-excel-v1'
NOTICE = 'Invoice export — not submitted to an accounting system'
COLUMNS = (
    ('invoice_id', 'Invoice ID', 39), ('review_revision', 'Review revision', 18),
    ('supplier_name', 'Supplier name', 32), ('supplier_address', 'Supplier address', 48),
    ('supplier_tax_id', 'Supplier tax ID', 24), ('invoice_number', 'Invoice number', 26),
    ('invoice_date', 'Invoice date', 16), ('due_date', 'Due date', 16), ('currency', 'Currency', 13),
    ('net_amount', 'Net amount', 20), ('tax_amount', 'Tax amount', 20), ('total_amount', 'Total amount', 20),
    ('discount_amount', 'Discount amount', 20), ('shipping_amount', 'Shipping amount', 20),
    ('approved_at', 'Approved at (UTC)', 35), ('export_id', 'Export ID', 39),
    ('mapping_version', 'Mapping version', 24), ('approval_method', 'Approval method', 22),
    ('policy_version', 'Policy version', 32),
)


def values_for(fields, task, export):
    if export['mapping_version'] != MAPPING_VERSION:
        raise ValueError('UNSUPPORTED_MAPPING_VERSION')
    values = {**fields, 'invoice_id': export['invoice_id'], 'review_revision': export['review_revision'],
              'export_id': export['export_id'], 'mapping_version': export['mapping_version'],
              'approved_at': task['decided_at'], 'approval_method': task['approval_method'],
              'policy_version': task['policy_version']}
    precision = CURRENCY_PRECISION[fields['currency']]
    unit = Decimal(1).scaleb(-precision)
    for key in MONEY_FIELDS:
        value = values.get(key)
        if value is not None:
            amount = Decimal(value)
            # Excel supports only 15 significant digits; never silently round the approved amount.
            if not amount.is_finite() or amount != amount.quantize(unit) or len(amount.normalize().as_tuple().digits) > 15 or abs(amount) >= Decimal('1e15'):
                raise ValueError('UNREPRESENTABLE_EXCEL_AMOUNT')
            values[key] = amount
    for key in ('invoice_date', 'due_date'):
        values[key] = date.fromisoformat(values[key]) if values.get(key) else None
    return values, '#,##0' + ('.' + '0' * precision if precision else '') + f' "{fields["currency"]}"'


def create_xlsx(fields, task, export):
    return create_rows_xlsx([(fields, task, export)])


def create_rows_xlsx(records):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = 'Invoice export'
    for index, (key, label, width) in enumerate(COLUMNS, 1):
        header = sheet.cell(1, index, label)
        header.font = Font(name='Calibri', bold=True, color='FFFFFF')
        header.fill = PatternFill('solid', fgColor='183044')
        header.alignment = Alignment(vertical='center', wrap_text=True)
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row_number, (fields, task, export) in enumerate(records, 2):
        values, money_format = values_for(fields, task, export)
        for index, (key, _, _) in enumerate(COLUMNS, 1):
            cell = sheet.cell(row_number, index, values.get(key))
            cell.font = Font(name='Calibri', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if isinstance(values.get(key), str):
                cell.data_type = 's'  # Including strings beginning with =, +, -, @ and leading zeros.
                cell.number_format = '@'
            elif key in MONEY_FIELDS:
                cell.number_format = money_format
            elif key in ('invoice_date', 'due_date'):
                cell.number_format = 'yyyy-mm-dd'
        sheet.row_dimensions[row_number].height = 64
    sheet.freeze_panes = 'A2'
    sheet.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNS))}{len(records) + 1}'
    sheet.row_dimensions[1].height = 32
    notice_row = len(records) + 3
    sheet.merge_cells(f'A{notice_row}:F{notice_row}')
    sheet.cell(notice_row, 1, NOTICE).font = Font(name='Calibri', italic=True, color='586979', size=12)
    sheet.row_dimensions[notice_row].height = 25
    fixed_time = datetime(2000, 1, 1)
    workbook.properties.creator = 'AP invoice demo'
    workbook.properties.created = fixed_time
    workbook.properties.modified = fixed_time
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    # Stable ZIP timestamps and core metadata make recovery of the same snapshot byte-identical.
    output = BytesIO()
    with ZipFile(buffer) as source, ZipFile(output, 'w', compression=ZIP_DEFLATED) as target:
        for name in sorted(source.namelist()):
            data = source.read(name)
            if name == 'docProps/core.xml':
                import re
                data = re.sub(rb'(<dcterms:modified[^>]*>).*?(</dcterms:modified>)', rb'\g<1>2000-01-01T00:00:00Z\2', data)
            entry = ZipInfo(name, (2000, 1, 1, 0, 0, 0))
            entry.create_system = 3  # Match the deployed Linux writer even during Windows recovery.
            entry.compress_type = ZIP_DEFLATED
            target.writestr(entry, data)
    return output.getvalue()


def verify_xlsx(data, fields, task, export):
    return verify_rows_xlsx(data, [(fields, task, export)])


def verify_rows_xlsx(data, records):
    with ZipFile(BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 16_000_000:
            raise ValueError('INVALID_WORKBOOK')
        if any('vbaproject' in n.lower() or 'externallinks/' in n.lower() for n in archive.namelist()):
            raise ValueError('ACTIVE_WORKBOOK_CONTENT')
        if any(b'TargetMode="External"' in archive.read(n) for n in archive.namelist() if n.endswith('.rels')):
            raise ValueError('EXTERNAL_WORKBOOK_LINK')
    workbook = load_workbook(BytesIO(data), data_only=False, keep_links=True)
    try:
        if workbook.sheetnames != ['Invoice export']:
            raise ValueError('INVALID_WORKBOOK_SHEETS')
        sheet = workbook.active
        if (sheet.freeze_panes != 'A2' or sheet.auto_filter.ref != f'A1:{get_column_letter(len(COLUMNS))}{len(records) + 1}'
                or sheet.cell(len(records) + 3, 1).value != NOTICE or sheet.max_row != len(records) + 3 or sheet.max_column != len(COLUMNS)):
            raise ValueError('INVALID_WORKBOOK_LAYOUT')
        for row in sheet:
            if any(cell.data_type == 'f' or cell.hyperlink for cell in row):
                raise ValueError('ACTIVE_WORKBOOK_CONTENT')
        for row_number, (fields, task, export) in enumerate(records, 2):
            values, money_format = values_for(fields, task, export)
            for index, (key, label, _) in enumerate(COLUMNS, 1):
                cell, expected = sheet.cell(row_number, index), values.get(key)
                actual = cell.value
                if sheet.cell(1, index).value != label:
                    raise ValueError('WORKBOOK_HEADER_MISMATCH')
                if isinstance(expected, Decimal):
                    valid = cell.data_type == 'n' and actual is not None and Decimal(str(actual)) == expected and cell.number_format == money_format
                elif isinstance(expected, date):
                    valid = isinstance(actual, datetime) and actual.date() == expected and cell.number_format == 'yyyy-mm-dd'
                else:
                    valid = actual == expected and (not isinstance(expected, str) or cell.data_type == 's')
                if not valid:
                    raise ValueError('WORKBOOK_SNAPSHOT_MISMATCH')
    finally:
        workbook.close()
    return hashlib.sha256(data).hexdigest()
