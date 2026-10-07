"""Offline Excel acceptance, concurrency, snapshot isolation and crash recovery."""
from concurrent.futures import ThreadPoolExecutor
from contextlib import closing
from datetime import datetime
from decimal import Decimal
from io import BytesIO
import json
import sqlite3
import subprocess
import sys
from uuid import uuid4

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from backend.exports import Exports
from backend.invoice_xlsx import COLUMNS, NOTICE
from backend.reviews import Reviews
from backend.test_review_decisions import env, seed, save, decide, state, HEADERS


def approved(env, automatic=False, **kwargs):
    invoice_id, attempt = seed(env, **kwargs)
    if automatic:
        result = Reviews(env[1]).automatic(invoice_id)
    else:
        assert save(env, invoice_id).status_code == 200
        result = decide(env, invoice_id).json()
    return invoice_id, attempt, result['event']['event_id']


def export(env, event_id, **extra):
    return env[0].post('/api/ap/exports', json={'event_id': event_id, **extra}, headers=HEADERS)


@pytest.mark.parametrize('automatic', [True, False])
def test_approved_export_roundtrip_and_download_is_read_only(env, automatic):
    invoice_id, _, event = approved(env, automatic=automatic, changes={'supplier_tax_id': '00123'}, number='000123')
    response = export(env, event)
    assert response.status_code == 200, response.text
    record = response.json()
    data = env[0].get(record['download_url'])
    assert data.status_code == 200 and data.content.startswith(b'PK')
    workbook = load_workbook(BytesIO(data.content))
    sheet = workbook.active
    cells = {key: sheet.cell(2, i) for i, (key, _, _) in enumerate(COLUMNS, 1)}
    assert cells['invoice_number'].value == '000123' and cells['invoice_number'].data_type == 's'
    assert cells['supplier_tax_id'].value == '00123'
    assert cells['total_amount'].data_type == 'n' and Decimal(str(cells['total_amount'].value)) == Decimal('119.00')
    assert cells['discount_amount'].value is None and cells['shipping_amount'].value is None
    assert cells['invoice_date'].value == datetime(2026, 10, 1)
    assert cells['approval_method'].value == ('automatic' if automatic else 'manual')
    assert cells['policy_version'].value == ('eur-gross-under-10000-v1' if automatic else None)
    assert sheet['A4'].value == NOTICE and sheet.freeze_panes == 'A2'
    workbook.close()
    assert export(env, event).json()['file_sha256'] == record['file_sha256']
    assert state(env, invoice_id)['status'] == 'approved'
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*),MAX(attempts) FROM exports').fetchone()[:] == (1, 1)
        assert db.execute('SELECT COUNT(*) FROM result_notifications').fetchone()[0] == 0


@pytest.mark.parametrize('prefix', ['=', '+', '-', '@'])
def test_external_text_never_becomes_formula(env, prefix):
    _, _, event = approved(env, changes={'supplier_name': prefix + 'SUM(1,2)', 'supplier_address': prefix + 'https://example.invalid'})
    record = export(env, event).json()
    workbook = load_workbook(BytesIO(env[0].get(record['download_url']).content))
    assert workbook.active['C2'].value == prefix + 'SUM(1,2)'
    assert workbook.active['C2'].data_type == 's'
    assert not any(c.data_type == 'f' or c.hyperlink for row in workbook.active for c in row)
    assert workbook._external_links == []
    workbook.close()


def test_snapshot_survives_machine_changes_and_mapping_deployment(env, monkeypatch):
    invoice_id, attempt = seed(env)
    save(env, invoice_id, changes={'supplier_address': 'Human corrected address'})
    event = decide(env, invoice_id).json()['event']['event_id']
    with closing(env[1]()) as db:
        db.execute('UPDATE processing_attempts SET result_json=? WHERE attempt_id=?', (json.dumps({'fields': {'total_amount': '9999', 'supplier_address': 'Wrong new draft'}}), attempt))
        db.commit()
    first = export(env, event).json()
    workbook = load_workbook(BytesIO(env[0].get(first['download_url']).content))
    assert workbook.active['D2'].value == 'Human corrected address'
    assert workbook.active['L2'].value == 119
    workbook.close()
    monkeypatch.setattr('backend.exports.MAPPING_VERSION', 'future-mapping')
    second = export(env, event).json()
    assert second['export_id'] == first['export_id'] and second['mapping_version'] == 'invoice-excel-v1'


def test_parallel_requests_share_one_identity_and_file(env):
    _, _, event = approved(env)
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda _: export(env, event), range(4)))
    assert all(r.status_code in (200, 409) for r in responses)
    final = export(env, event).json()
    assert all(r.json()['export_id'] == final['export_id'] for r in responses if r.status_code == 200)
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*) FROM exports').fetchone()[0] == 1
    assert len(list((env[2] / 'exports').glob('*.xlsx'))) == 1


def test_forged_pending_rejected_and_stale_events_cannot_export(env):
    assert export(env, str(uuid4())).status_code == 404
    invoice_id, _ = seed(env)
    assert export(env, invoice_id).status_code == 404
    save(env, invoice_id)
    event = decide(env, invoice_id, 'reject').json()['event']['event_id']
    assert export(env, event).status_code == 409
    _, _, approved_event = approved(env)
    with closing(env[1]()) as db:
        db.execute('UPDATE review_tasks SET current_revision=current_revision+1 WHERE status=\'approved\'')
        db.commit()
    assert export(env, approved_event).status_code == 409
    with closing(env[1]()) as db:
        assert db.execute('SELECT COUNT(*) FROM exports').fetchone()[0] == 0


def test_export_write_auth_and_payload_are_strict(env, monkeypatch):
    _, _, event = approved(env)
    assert env[0].post('/api/ap/exports', json={'event_id': event}).status_code == 403
    assert export(env, event, file_path='../outside.xlsx').status_code == 422
    assert env[0].post('/api/ap/exports', json={'event_id': event}, headers={**HEADERS, 'Origin': 'http://evil.invalid'}).status_code == 403
    monkeypatch.setenv('AP_REVIEW_EVENT_TOKEN', 'isolated-token')
    assert env[0].post('/api/ap/exports', json={'event_id': event}, headers={'X-AP-Event-Token': 'wrong'}).status_code == 403
    assert env[0].post('/api/ap/exports', json={'event_id': event}, headers={'X-AP-Event-Token': 'isolated-token'}).status_code == 200


def test_write_failure_is_bounded_and_does_not_undo_approval(env, monkeypatch):
    invoice_id, _, event = approved(env)
    def fail(*args):
        raise PermissionError('Private path must not appear in response')
    monkeypatch.setattr(Exports, 'write_file', fail)
    ids = []
    for _ in range(5):
        response = export(env, event)
        assert response.status_code == 503 and 'Private path' not in response.text
        ids.append(response.json()['detail']['export_id'])
    assert len(set(ids)) == 1 and export(env, event).status_code == 409
    assert state(env, invoice_id)['status'] == 'approved'


def test_file_saved_but_database_commit_fails_reuses_orphan(env, monkeypatch):
    _, _, event = approved(env)
    original = Exports.complete
    def failed_commit(*args):
        raise sqlite3.OperationalError('injected commit failure')
    monkeypatch.setattr(Exports, 'complete', failed_commit)
    failed = export(env, event)
    assert failed.status_code == 503
    export_id = failed.json()['detail']['export_id']
    path = env[2] / 'exports' / (export_id + '.xlsx')
    before = path.read_bytes()
    monkeypatch.setattr(Exports, 'complete', original)
    monkeypatch.setattr(Exports, 'write_file', lambda *args: pytest.fail('Valid orphan should be reused'))
    recovered = export(env, event).json()
    assert recovered['export_id'] == export_id and recovered['attempts'] == 2
    assert path.read_bytes() == before


@pytest.mark.parametrize('damage', ['missing', 'corrupt'])
def test_completed_file_recovery_is_explicit_and_same_hash(env, damage):
    _, _, event = approved(env)
    first = export(env, event).json()
    path = env[2] / first['file_path']
    if damage == 'missing':
        path.unlink()
    else:
        path.write_bytes(b'corrupt workbook')
    assert env[0].get(first['download_url']).status_code == 409
    assert env[0].get('/api/ap/exports/' + first['export_id']).json()['file_health'] == 'missing_or_invalid'
    with closing(env[1]()) as db:
        assert db.execute('SELECT attempts FROM exports').fetchone()[0] == 1
    recovered = export(env, event).json()
    assert recovered['export_id'] == first['export_id'] and recovered['file_sha256'] == first['file_sha256']


def test_crash_claim_survives_and_expired_lease_retries(env, monkeypatch):
    _, _, event = approved(env)
    original = Exports.write_file
    def crash(*args):
        raise SystemExit('Simulated abrupt worker interruption')
    monkeypatch.setattr(Exports, 'write_file', crash)
    with pytest.raises(SystemExit):
        Exports(env[1], env[2]).create(event)
    assert export(env, event).status_code == 409
    with closing(env[1]()) as db:
        assert db.execute('SELECT status,attempts FROM exports').fetchone()[:] == ('writing', 1)
        db.execute("UPDATE exports SET updated_at='2000-01-01T00:00:00+00:00'")
        db.commit()
    monkeypatch.setattr(Exports, 'write_file', original)
    assert export(env, event).json()['attempts'] == 2


@pytest.mark.parametrize('currency,total,format_part', [('JPY', '119', '#,##0 "JPY"'), ('KWD', '1.25', '#,##0.000 "KWD"')])
def test_currency_precision(env, currency, total, format_part):
    _, _, event = approved(env, total=total, changes={'currency': currency})
    record = export(env, event).json()
    workbook = load_workbook(BytesIO(env[0].get(record['download_url']).content))
    assert workbook.active['L2'].number_format == format_part
    workbook.close()


def test_unrepresentable_excel_amount_is_not_silently_rounded(env):
    _, _, event = approved(env, total='123456789012345.67')
    result = export(env, event)
    assert result.status_code == 503 and result.json()['detail']['error'] == 'UNREPRESENTABLE_EXCEL_AMOUNT'


@pytest.mark.parametrize('phase', ['partial', 'renamed'])
def test_real_process_exit_recovers_same_export(env, phase):
    _, _, event = approved(env)
    script = '''
import os,sqlite3,sys
from pathlib import Path
from backend.exports import Exports
root=Path(sys.argv[1])
def connect():
 db=sqlite3.connect(root/'invoices.sqlite3',timeout=30);db.row_factory=sqlite3.Row;db.execute('PRAGMA foreign_keys=ON');return db
class CrashExport(Exports):
 def write_file(self,path,data):
  if sys.argv[3]=='partial':
   path.parent.mkdir(exist_ok=True);path.with_suffix('.tmp').write_bytes(data[:50])
  else:
   super().write_file(path,data)
  os._exit(77)
CrashExport(connect,root).create(sys.argv[2])
'''
    process = subprocess.run([sys.executable, '-c', script, str(env[2]), event, phase], capture_output=True, timeout=30)
    assert process.returncode == 77, process.stderr
    with closing(env[1]()) as db:
        before = dict(db.execute('SELECT * FROM exports').fetchone())
        assert before['status'] == 'writing' and before['attempts'] == 1
        db.execute("UPDATE exports SET updated_at='2000-01-01T00:00:00+00:00'")
        db.commit()
    recovered = export(env, event).json()
    assert recovered['export_id'] == before['export_id'] and recovered['status'] == 'completed'
    assert not list((env[2] / 'exports').glob('*.tmp'))


def test_notification_failure_never_recreates_export(env):
    _, _, event = approved(env)
    first = export(env, event).json()
    with closing(env[1]()) as db:
        db.execute("INSERT INTO result_notifications(notification_id,event_id,export_id,category,status) VALUES(?,?,?,'export_completed','failed')", (str(uuid4()), event, first['export_id']))
        db.commit()
    assert export(env, event).json() == first


def test_stored_path_cannot_escape_controlled_directory(env):
    _, _, event = approved(env)
    first = export(env, event).json()
    with closing(env[1]()) as db:
        db.execute("UPDATE exports SET file_path='../outside.xlsx'")
        db.commit()
    assert env[0].get(first['download_url']).status_code == 409
    assert export(env, event).status_code == 503
    assert not (env[2] / 'outside.xlsx').exists()


def test_writer_hash_does_not_depend_on_zip_host_platform(env, monkeypatch):
    import backend.invoice_xlsx as writer
    _, _, event = approved(env)
    record = export(env, event).json()
    path = env[2] / record['file_path']
    original = writer.ZipInfo
    def windows_entry(*args, **kwargs):
        info = original(*args, **kwargs)
        info.create_system = 0
        return info
    monkeypatch.setattr(writer, 'ZipInfo', windows_entry)
    path.unlink()
    recovered = export(env, event).json()
    assert recovered['file_sha256'] == record['file_sha256']
